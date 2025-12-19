"""Generate an editable PPT slide for the high-quality data platform matrix.

This script recreates the content from the provided reference slide using
`python-pptx`. All text is editable after generation. Colors and spacing are
kept simple so the resulting file is easy to adjust.

Usage:
    pip install python-pptx
    python script/generate_high_quality_data_ppt.py output.pptx

If no output path is provided, the script writes `high_quality_data_platform.pptx`
into the current working directory.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Iterable, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SlideBounds = Tuple[float, float, float, float]


def _add_textbox(slide, bounds: SlideBounds, text: str, *, font_size=12, bold=False, color=RGBColor(0, 0, 0), align=PP_ALIGN.LEFT):
    """Add a textbox with standard padding and styling."""
    left, top, width, height = (Inches(v) for v in bounds)
    shape = slide.shapes.add_textbox(left, top, width, height)
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True

    for paragraph in text_frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def _add_heading(slide, text: str):
    _add_textbox(
        slide,
        bounds=(0.5, 0.3, 11.5, 0.6),
        text=text,
        font_size=22,
        bold=True,
        color=RGBColor(0, 96, 170),
    )


def _add_subtitle(slide, text: str):
    _add_textbox(
        slide,
        bounds=(0.5, 1.0, 11.5, 0.7),
        text=text,
        font_size=14,
        color=RGBColor(64, 64, 64),
    )


def _add_labeled_box(slide, title: str, content: Iterable[str], bounds: SlideBounds, *, title_bg=RGBColor(0, 96, 170)):
    """Create a boxed section with a colored header and bullet-like lines."""
    left, top, width, height = (Inches(v) for v in bounds)
    # Outer box
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    box.line.color.rgb = RGBColor(191, 191, 191)
    box.fill.fore_color.rgb = RGBColor(242, 242, 242)

    # Title bar
    title_height = Inches(0.5)
    title_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, title_height)
    title_box.fill.solid()
    title_box.fill.fore_color.rgb = title_bg
    title_box.line.fill.background()

    title_frame = title_box.text_frame
    title_frame.text = title
    for paragraph in title_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Content
    content_text = "\n".join(f"• {line}" for line in content)
    top_in = bounds[1] + 0.5  # leave space for the title bar
    _add_textbox(
        slide,
        bounds=(bounds[0] + 0.2, top_in, bounds[2] - 0.4, bounds[3] - 0.6),
        text=content_text,
        font_size=12,
    )



def build_slide(output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    _add_heading(slide, "1.1 高质量数据供给与管理平台：构建高质量数据集平台矩阵")
    _add_subtitle(
        slide,
        "构建统一的高质量数据供给生产与供给全链条协同的平台矩阵。驱动各场景的数字化治理。覆盖多样化场景的数据供给需求，行业、场景真实数据与仿真数据。")

    # Left column: 城市大脑工程
    left_box_text = [
        "一网统管",
        "数字政府",
        "智慧能源",
        "智慧交通",
        "AI+交通",
        "智慧工业园区",
        "智慧建筑工程",
        "智慧应急管理",
        "社会治理",
        "智慧养老",
        "智慧环保",
        "智慧医保",
        "智慧政法",
    ]
    _add_labeled_box(
        slide,
        title="智慧城市大脑工程",
        content=left_box_text,
        bounds=(0.5, 1.8, 3.0, 5.3),
        title_bg=RGBColor(0, 122, 204),
    )

    # Center description
    _add_textbox(
        slide,
        bounds=(3.7, 1.5, 6.2, 0.7),
        text=(
            "平台：构建统一的高质量数据供给生产与供给全链条协同的平台矩阵，驱动各场景的数字化治理。\n"
            "数据：覆盖多样化场景的数据供给需求，行业、场景真实数据与仿真数据。"
        ),
        font_size=13,
        color=RGBColor(64, 64, 64),
    )

    # MaaS 多模型协同
    maas_left = [
        "高质量数据生产&融合（知识图谱、智能数据标注、数据增强）",
        "知识图谱：行业知识图谱自动构建、数据辅助采集、公共环境补全",
        "智能数据标注：数据标注过程自动化、零样本、几样本标注",
        "数据增强：AI生产、自动清洗、自动修补",
        "多模态数据融合（打通不同要素数据，实现多模态融合集成）",
        "结构化数据：场景模型治理为基础，构建行业知识图谱，数据采集引导知识图谱修补",
        "非结构化数据：语言大模型、视觉模型、语音模型、数学模型、心智模型",
        "多模态数据：利用数据挖掘构建多模态统一数据模型，实现多模态模型协同",
    ]

    _add_labeled_box(
        slide,
        title="MaaS 多模型协同",
        content=maas_left,
        bounds=(3.7, 2.3, 4.0, 4.8),
        title_bg=RGBColor(76, 114, 176),
    )

    # Intelligent automated data engineering platform
    right_box_text = [
        "数据集生产：数据采集系统自动化，数据标注过程自动化，标准化数据集自动化生产，人工介入闭环",
        "模型研究生产与微调：模型训练平台，模型自动化管理，多模态模型微调，模拟场景构造，模型迭代与数据重塑",
        "数据存储与数据供给：数据全生命周期治理，数据湖与知识图谱，向量数据库，数据服务化为下游供给多模态模型应用，智慧城市大脑应用",
        "安全与治理：安全体系搭建，模型安全与数据安全，隐私安全与合规，模型可解释与自动化，模型与数据审计",
    ]
    _add_labeled_box(
        slide,
        title="智能全自动数据工程平台",
        content=right_box_text,
        bounds=(7.9, 2.3, 4.8, 4.8),
        title_bg=RGBColor(0, 153, 255),
    )

    # Bottom note
    _add_textbox(
        slide,
        bounds=(3.7, 7.2, 8.8, 0.3),
        text="平台服务：统一数据与业务数据产品，封装在标准引擎工具中，供各场景复用 | PaaS模式的标准化引擎与工具",
        font_size=11,
        color=RGBColor(96, 96, 96),
        align=PP_ALIGN.RIGHT,
    )

    prs.save(output_path)
    return output_path


def main():
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("high_quality_data_platform.pptx")
    path = build_slide(output)
    print(f"PPT saved to: {path}")


if __name__ == "__main__":
    main()
